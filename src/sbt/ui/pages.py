"""Side-by-side page images for Review: the original page and the translated page, with the paragraphs that
have quality findings outlined.

PDF pages are drawn by PyMuPDF. PowerPoint slides are drawn by the PowerPoint installed on this PC (opened
read-only and without a window; a PowerPoint the user already has open is left alone). Images are kept in a
cache folder in the user's data folder and deleted when the next job starts or the app closes.
"""
from __future__ import annotations

import base64
import hashlib
import re
import shutil
import subprocess
from pathlib import Path

from sbt import settings as settings_mod
from sbt.domain.models import DocumentModel, Severity
from sbt.quality.findings import QcReport

_PS = r"""
param([string]$Pptx, [string]$OutDir)
$ErrorActionPreference = "Stop"
New-Item -ItemType Directory -Force $OutDir | Out-Null
$wasRunning = [bool](Get-Process POWERPNT -ErrorAction SilentlyContinue)
$app = New-Object -ComObject PowerPoint.Application
try {
    $pres = $app.Presentations.Open($Pptx, -1, 0, 0)
    try {
        foreach ($slide in $pres.Slides) {
            $slide.Export((Join-Path $OutDir ("slide{0:D3}.png" -f $slide.SlideIndex)), "PNG", 1280, 720)
        }
    } finally { $pres.Close() }
} finally {
    if (-not $wasRunning -and $app.Presentations.Count -eq 0) { $app.Quit() }
    [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($app)
}
"""
_SHAPE = re.compile(r"^s(\d+)/(\d+)")


def cache_dir() -> Path:
    return settings_mod.data_dir() / "cache" / "pages"


def clear_cache() -> None:
    shutil.rmtree(cache_dir(), ignore_errors=True)


def powerpoint_available() -> bool:
    try:
        import winreg
        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, r"PowerPoint.Application\CLSID"))
        return True
    except OSError:
        return False


def _png(path: Path) -> str:
    return "data:image/png;base64," + base64.b64encode(path.read_bytes()).decode("ascii")


def _slides(pptx: Path) -> Path:
    """Folder with slideNNN.png for this file version (exported once)."""
    key = hashlib.sha256(f"{pptx.resolve()}|{pptx.stat().st_mtime_ns}".encode()).hexdigest()[:16]
    out = cache_dir() / key
    if not any(out.glob("slide*.png")):
        out.mkdir(parents=True, exist_ok=True)
        script = cache_dir() / "export_slides.ps1"
        script.write_text(_PS, encoding="utf-8")
        subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-File",
                        str(script), "-Pptx", str(pptx.resolve()), "-OutDir", str(out)],
                       check=True, capture_output=True, timeout=600,
                       creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    return out


def _pdf_page(path: Path, index: int) -> str:
    import pymupdf
    with pymupdf.open(path) as doc:
        page = doc[index]
        pix = page.get_pixmap(matrix=pymupdf.Matrix(1100 / page.rect.width, 1100 / page.rect.width), alpha=False)
        return "data:image/png;base64," + base64.b64encode(pix.tobytes("png")).decode("ascii")


def _boxes(doc: DocumentModel, report: QcReport, page_no: int, original: Path) -> list[dict[str, object]]:
    """Outlines (0–1 coordinates) of the paragraphs on this page that have findings."""
    by_seg = report.by_segment()
    flagged = {sid: fs for sid, fs in by_seg.items() if any(f.severity != Severity.INFO for f in fs)}
    out: list[dict[str, object]] = []
    if original.suffix.lower() == ".pdf":
        import pymupdf
        with pymupdf.open(original) as pdf:
            page_rect = pdf[page_no - 1].rect
        for s in doc.segments:
            if s.id in flagged and s.container == page_no and "rect" in s.geometry:
                x0, y0, x1, y1 = s.geometry["rect"]  # type: ignore[misc]
                out.append(_box(s.id, x0 / page_rect.width, y0 / page_rect.height, (x1 - x0) / page_rect.width,
                                (y1 - y0) / page_rect.height, flagged[s.id]))
        return out
    from pptx import Presentation
    prs = Presentation(str(original))
    slide = prs.slides[page_no - 1]
    shapes = {sh.shape_id: sh for sh in slide.shapes}
    seen: set[int] = set()
    for s in doc.segments:
        m = _SHAPE.match(s.id)
        if s.id not in flagged or s.container != page_no or not m or s.kind.value == "note":
            continue
        sh = shapes.get(int(m.group(2)))
        if sh is None or sh.width is None or int(m.group(2)) in seen:
            continue
        seen.add(int(m.group(2)))
        out.append(_box(s.id, sh.left / prs.slide_width, sh.top / prs.slide_height, sh.width / prs.slide_width,
                        sh.height / prs.slide_height, flagged[s.id]))
    return out


def _box(sid: str, x: float, y: float, w: float, h: float, findings: list) -> dict[str, object]:  # type: ignore[type-arg]
    worst = "error" if any(f.severity == Severity.ERROR for f in findings) else "warning"
    return {"segment": sid, "x": x, "y": y, "w": w, "h": h, "severity": worst,
            "text": "; ".join(dict.fromkeys(f.describe() for f in findings if f.severity != Severity.INFO))}


def page_view(doc: DocumentModel, report: QcReport, original: Path, translated: Path, page_no: int
              ) -> dict[str, object]:
    total = doc.container_count
    page_no = max(1, min(int(page_no), total))
    base: dict[str, object] = {"page": page_no, "pages": total,
                               "unit": "page" if original.suffix.lower() == ".pdf" else "slide"}
    if not translated.exists():
        return {**base, "available": False, "reason": "The translated file was not found."}
    if original.suffix.lower() == ".pdf":
        return {**base, "available": True, "original": _pdf_page(original, page_no - 1),
                "translated": _pdf_page(translated, page_no - 1),
                "boxes": _boxes(doc, report, page_no, original)}
    if not powerpoint_available():
        return {**base, "available": False,
                "reason": "Slide pictures need Microsoft PowerPoint on this PC. The text comparison below works without it."}
    try:
        a, b = _slides(original), _slides(translated)
    except (subprocess.SubprocessError, OSError):
        return {**base, "available": False, "reason": "PowerPoint could not draw the slides."}
    name = f"slide{page_no:03d}.png"
    if not (a / name).exists() or not (b / name).exists():
        return {**base, "available": False, "reason": "PowerPoint could not draw this slide."}
    return {**base, "available": True, "original": _png(a / name), "translated": _png(b / name),
            "boxes": _boxes(doc, report, page_no, original)}

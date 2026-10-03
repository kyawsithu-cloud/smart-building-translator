"""Approximate text layout for overflow detection and gentle font reduction.

python-pptx has no layout engine, so we estimate: resolve each paragraph's effective font size through the
inheritance chain (run → shape list style → layout/master placeholder → master text styles), measure with a
real Windows font via Pillow, wrap greedily, and compare the resulting height to the box. It is an estimate;
the report says so.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

from lxml import etree
from PIL import ImageFont

from sbt.languages import Language
from sbt import winfonts
from sbt.parsers.pptx_walk import A

P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
EMU_PER_PT = 12700
DEFAULT_SZ = 18.0
LINE_FACTOR = 1.2
TOLERANCE = 2.0            # pt: PowerPoint lets text run into the box's inner margin without visible overflow
FONTS = Path(r"C:\Windows\Fonts")
_FONT_FILES = {
    "ja": ["meiryo.ttc", "YuGothM.ttc", "msgothic.ttc"],
    "zh": ["msyh.ttc", "simsun.ttc"],
    "ko": ["malgun.ttf"],
    "th": ["LeelawUI.ttf", "tahoma.ttf"],
    "my": ["mmrtext.ttf"],
    "*": ["calibri.ttf", "arial.ttf"],
}
_TOKEN = re.compile(r"[A-Za-z0-9\u00C0-\u024F\-./:@_%\u00b0\u00b1+#()]+\s*|\s+|.", re.S)
_SPACED = re.compile(r"\S+\s*|\s+")


@dataclass
class FitResult:
    scale: float
    overflow: bool                 # the translation makes the text exceed its box
    grows: bool = False            # auto-growing box: it will get taller (may cover other content)
    inherited: bool = False        # exceeds its box, but the original text did too


@lru_cache(maxsize=8)
def _font(lang_code: str) -> ImageFont.FreeTypeFont | None:
    for name in _FONT_FILES.get(lang_code, []) + _FONT_FILES["*"]:
        path = FONTS / name
        if path.exists():
            return ImageFont.truetype(str(path), 100)
    return None


Fonts = tuple[ImageFont.FreeTypeFont | None, ImageFont.FreeTypeFont | None]   # (Latin text, native script)


@lru_cache(maxsize=64)
def _family_font(family: str) -> ImageFont.FreeTypeFont | None:
    found = winfonts.face(family) if family else None
    if found is None:
        return None
    try:
        return ImageFont.truetype(str(found[0]), 100, index=found[1])
    except OSError:
        return None



def _shape_fonts(shape, lang: Language) -> Fonts:
    """The fonts PowerPoint will use in this text box: the first run's own typefaces, else the theme's
    (major fonts for titles). Japanese/Chinese/Korean use the East Asian slot, Thai/Burmese the complex-script
    slot. Narrow theme fonts (e.g. Cordia New for Thai) wrap very differently from a stand-in font."""
    try:
        part = shape.part.slide.slide_layout.slide_master.part
        cache = part.__dict__.setdefault("_sbt_theme_fonts", {})     # lives as long as this presentation
        if lang.code not in cache:
            from pptx.opc.constants import RELATIONSHIP_TYPE as RT
            cache[lang.code] = winfonts.theme_fonts(etree.fromstring(part.part_related_by(RT.THEME).blob),
                                                    lang.code)
        theme = cache[lang.code]
    except (AttributeError, KeyError, etree.XMLSyntaxError):
        theme = {}
    level = "mj" if shape.is_placeholder and str(shape.placeholder_format.type).split(" ")[0].upper() in (
        "TITLE", "CENTER_TITLE", "VERTICAL_TITLE") else "mn"
    rpr = shape.text_frame._txBody.find(f".//{A}r/{A}rPr")
    slot = "cs" if lang.code in ("th", "my") else "ea"

    def typeface(child: str, kind: str) -> str:
        el = rpr.find(f"{A}{child}") if rpr is not None else None
        face = el.get("typeface", "") if el is not None else f"+{level}-{kind}"
        return theme.get(face, "") if face.startswith("+") else face

    latin = _family_font(typeface("latin", "lt")) or _font("*")
    native = _family_font(typeface(slot, slot)) or _font(lang.code)
    return latin, native


def _text_width_pt(text: str, size_pt: float, lang: Language, fonts: Fonts | None = None) -> float:
    # PowerPoint renders Latin characters with the Latin theme font even in Japanese text.
    native = lang.script is not None and lang.code != "en" and lang.script.search(text) and not text.isascii()
    if fonts is not None:
        font = fonts[1] if native else fonts[0]
    else:
        font = _font(lang.code if native else "*")
    if font is None:
        return len(text) * size_pt * (1.0 if lang.east_asian else 0.5)
    return font.getlength(text) * size_pt / 100


def _lvl_sz(lst_style: etree._Element | None, lvl: int) -> float | None:
    if lst_style is None:
        return None
    d = lst_style.find(f"{A}lvl{lvl + 1}pPr/{A}defRPr")
    if d is not None and d.get("sz"):
        return int(d.get("sz")) / 100
    return None


def _style_chain(shape, lvl: int) -> list[etree._Element]:
    """List-style levels (lvlNpPr) that apply to paragraphs at this level, nearest first: the shape's own list
    style, its layout and master placeholders, then the master's title/body/other text style."""
    chain = []
    node = shape
    for _ in range(3):   # slide → layout → master
        body = node.element.find(f".//{P}txBody")
        lst = body.find(f"{A}lstStyle") if body is not None else None
        level = lst.find(f"{A}lvl{lvl + 1}pPr") if lst is not None else None
        if level is not None:
            chain.append(level)
        node = getattr(node, "_base_placeholder", None) if node.is_placeholder else None
        if node is None:
            break
    try:
        master = shape.part.slide.slide_layout.slide_master.element
    except AttributeError:
        return chain
    style = "otherStyle"
    if shape.is_placeholder:     # a subtitle uses the body style (its type name contains "TITLE" too)
        kind = str(shape.placeholder_format.type).split(" ")[0].upper()
        style = "titleStyle" if kind in ("TITLE", "CENTER_TITLE", "VERTICAL_TITLE") else "bodyStyle"
    level = master.find(f"{P}txStyles/{P}{style}/{A}lvl{lvl + 1}pPr")
    if level is not None:
        chain.append(level)
    return chain


def _inherited_sz(shape, lvl: int) -> float:
    """Walk placeholder inheritance to find the default size for paragraphs at this level."""
    for level in _style_chain(shape, lvl):
        d = level.find(f"{A}defRPr")
        if d is not None and d.get("sz"):
            return int(d.get("sz")) / 100
    return DEFAULT_SZ


@dataclass
class _Para:
    text: str
    size: float                      # effective font size (pt)
    rprs: list[etree._Element]       # run properties to scale
    margin: float                    # left margin (pt): bullets and indents narrow the line
    before: tuple[str, float]        # ("pct", lines) or ("pts", points)
    after: tuple[str, float]
    line: tuple[str, float]          # line spacing: ("pct", multiplier) or ("pts", exact height)


def _spacing(el: etree._Element | None, default: tuple[str, float]) -> tuple[str, float]:
    if el is None:
        return default
    pct, pts = el.find(f"{A}spcPct"), el.find(f"{A}spcPts")
    if pct is not None and pct.get("val"):
        return ("pct", int(pct.get("val")) / 100000)
    if pts is not None and pts.get("val"):
        return ("pts", int(pts.get("val")) / 100)
    return default


def _paragraphs(shape) -> list[_Para]:
    """Each paragraph with the size and spacing PowerPoint will use (own settings, else inherited)."""
    out = []
    for p in shape.text_frame.paragraphs:
        rprs = [r for r in p._p.iter(f"{A}rPr", f"{A}endParaRPr")]
        explicit = [int(r.get("sz")) / 100 for r in rprs if r.get("sz")]
        size = max(explicit) if explicit else _inherited_sz(shape, p.level)
        own = p._p.find(f"{A}pPr")
        chain = ([own] if own is not None else []) + _style_chain(shape, p.level)
        margin = next((int(el.get("marL")) / EMU_PER_PT for el in chain if el.get("marL")), 0.0)
        text = "".join(t.text or "" for t in p._p.iter(f"{A}t"))
        out.append(_Para(text, size, rprs, margin, _spacing(_first(chain, "spcBef"), ("pts", 0.0)),
                         _spacing(_first(chain, "spcAft"), ("pts", 0.0)),
                         _spacing(_first(chain, "lnSpc"), ("pct", 1.0))))
    return out


def _first(chain: list[etree._Element], name: str) -> etree._Element | None:
    """The nearest definition of a paragraph property (the paragraph's own, else inherited)."""
    return next((c for c in (el.find(f"{A}{name}") for el in chain) if c is not None), None)


def _paragraph_sizes(shape) -> list[tuple[str, float, list[etree._Element]]]:
    """(text, effective size pt, rPr elements to scale) for each paragraph."""
    return [(p.text, p.size, p.rprs) for p in _paragraphs(shape)]


def _box(shape) -> tuple[float, float, bool]:
    body = shape.text_frame._txBody.find(f"{A}bodyPr")
    get = lambda k, d: int(body.get(k)) if body is not None and body.get(k) else d  # noqa: E731
    w = (shape.width - get("lIns", 91440) - get("rIns", 91440)) / EMU_PER_PT
    h = (shape.height - get("tIns", 45720) - get("bIns", 45720)) / EMU_PER_PT
    wrap = body is None or body.get("wrap") != "none"
    return w, h, wrap


def _lines(text: str, size: float, width: float, lang: Language, wrap: bool,
           fonts: Fonts | None = None) -> tuple[int, float]:
    """Returns (line count, widest line width)."""
    count, widest = 0, 0.0
    # Burmese has no line-break opportunities between syllables in PowerPoint: it wraps at spaces only
    # (measured against PowerPoint's layout); other scripts may break between characters.
    pattern = _SPACED if lang.code == "my" else _TOKEN
    for hard_line in text.split("\n") or [""]:
        count += 1
        line_w = 0.0
        for tok in pattern.findall(hard_line):
            w = _text_width_pt(tok, size, lang, fonts)
            if wrap and line_w + w > width and line_w > 0:
                count += 1
                widest = max(widest, line_w)
                line_w = _text_width_pt(tok.lstrip(), size, lang, fonts)
            else:
                line_w += w
            if wrap and width > 0 and line_w > width:          # a word longer than the line is broken
                extra = int(line_w // width)
                count += extra
                line_w -= extra * width
        widest = max(widest, line_w)
    return count, widest


def _space(spec: tuple[str, float], line_height: float) -> float:
    return spec[1] * line_height if spec[0] == "pct" else spec[1]


def estimate_height(shape, lang: Language, scale: float = 1.0) -> tuple[float, float]:
    """(needed height pt, box height pt). Line height is 1.2 × font size in every script (measured in
    PowerPoint), times the paragraph's line spacing; space before a paragraph is ignored for the first one."""
    width, height, wrap = _box(shape)
    total = 0.0
    paras = _paragraphs(shape)
    fonts = _shape_fonts(shape, lang)
    for i, p in enumerate(paras):
        size = p.size * scale
        line_h = size * LINE_FACTOR * p.line[1] if p.line[0] == "pct" else p.line[1]
        n, widest = _lines(p.text, size, max(width - p.margin, width * 0.3), lang, wrap, fonts)
        total += n * line_h
        if i > 0:
            total += _space(p.before, size * LINE_FACTOR)
        if i < len(paras) - 1:
            total += _space(p.after, size * LINE_FACTOR)
        if not wrap and widest > width:
            total = max(total, height * widest / width)   # treat horizontal overflow as overflow
    return total, height


def fit_text_frame(shape, lang: Language, min_scale: float, allowed_height: float | None = None) -> FitResult:
    """Shrink the font (not below min_scale) until the text fits. `allowed_height`: what the original text
    needed — a box that already overflowed is not "fixed", only kept from getting worse."""
    body = shape.text_frame._txBody.find(f"{A}bodyPr")
    needed, box_h = estimate_height(shape, lang)
    if body is not None and body.find(f"{A}spAutoFit") is not None:      # shape grows with its text
        return FitResult(1.0, False, grows=needed > box_h + TOLERANCE)
    limit = max(box_h, allowed_height or 0.0) + TOLERANCE
    if needed <= limit:
        return FitResult(1.0, False, inherited=needed > box_h + TOLERANCE)
    scale = 1.0
    while scale - 0.05 >= min_scale - 1e-9:
        scale = round(scale - 0.05, 2)
        if estimate_height(shape, lang, scale)[0] <= limit:
            break
    for p in shape.text_frame.paragraphs:
        inherited = _inherited_sz(shape, p.level)
        for rpr in p._p.iter(f"{A}rPr", f"{A}endParaRPr"):
            size = int(rpr.get("sz")) / 100 if rpr.get("sz") else inherited
            rpr.set("sz", str(int(round(size * scale * 100))))
    after = estimate_height(shape, lang, 1.0)[0]
    return FitResult(scale, after > limit, inherited=box_h + TOLERANCE < after <= limit)

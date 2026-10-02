"""Deterministic traversal of a presentation's text, shared by the PPTX parser and renderer.

The parser and renderer both open the ORIGINAL file and walk it in the same order, so a segment id such as
"s3/12/p0" always identifies the same paragraph. No state needs to be kept between the two.
"""
from __future__ import annotations

import copy
from collections.abc import Iterator
from dataclasses import dataclass

from lxml import etree
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.presentation import Presentation as PresentationT
from pptx.text.text import TextFrame

from sbt.domain.models import SegmentKind

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
# rPr attributes that do not affect appearance; ignored when grouping runs into formatting classes.
_VOLATILE_ATTRS = ("lang", "altLang", "dirty", "err", "noProof", "smtClean", "smtId", "bmk")


@dataclass
class ParagraphRef:
    segment_id: str
    kind: SegmentKind
    slide_no: int
    shape_key: str
    element: etree._Element           # a:p
    text_frame: TextFrame | None
    shape: object | None              # owning shape (None for table cells / notes: sized differently)


@dataclass
class EncodedParagraph:
    text: str                                       # tagged text
    class_rpr: dict[int | None, etree._Element | None]   # tag id -> rPr template (None key = dominant)
    has_field: bool


def _class_key(rpr: etree._Element | None) -> str:
    if rpr is None:
        return ""
    clean = copy.deepcopy(rpr)
    for attr in _VOLATILE_ATTRS:
        clean.attrib.pop(attr, None)
    return etree.tostring(clean, method="c14n").decode()


def encode_paragraph(p: etree._Element) -> EncodedParagraph:
    pieces: list[tuple[str, str]] = []           # (class key, text)
    rprs: dict[str, etree._Element | None] = {}
    has_field = False
    for child in p:
        tag = etree.QName(child).localname
        if tag in ("r", "fld"):
            has_field |= tag == "fld"
            rpr = child.find(f"{A}rPr")
            key = _class_key(rpr)
            rprs.setdefault(key, rpr)
            t = child.find(f"{A}t")
            pieces.append((key, (t.text or "") if t is not None else ""))
        elif tag == "br":
            rpr = child.find(f"{A}rPr")
            key = pieces[-1][0] if pieces else _class_key(rpr)
            rprs.setdefault(key, rpr)
            pieces.append((key, "\n"))

    # Dominant class = most characters; it stays untagged to keep prompts simple.
    totals: dict[str, int] = {}
    for key, text in pieces:
        totals[key] = totals.get(key, 0) + len(text.strip())
    dominant = max(totals, key=lambda k: totals[k]) if totals else ""
    ids: dict[str, int | None] = {dominant: None}
    for key, _ in pieces:
        if key not in ids:
            ids[key] = len(ids)          # 1, 2, … in order of first appearance

    out: list[str] = []
    for key, text in pieces:
        tid = ids[key]
        if tid is None or not text.strip():   # whitespace-only spans don't need their own tag
            out.append(text)
        else:
            out.append(f"<g{tid}>{text}</g{tid}>")
    merged = "".join(out)
    # merge adjacent identical tags produced by consecutive runs of the same class
    for tid in set(i for i in ids.values() if i is not None):
        merged = merged.replace(f"</g{tid}><g{tid}>", "")
    # .get: a paragraph can have no runs at all (empty line), so the dominant class may have no rPr
    return EncodedParagraph(merged, {tid: rprs.get(key) for key, tid in ids.items()}, has_field)


# --- charts, SmartArt, pictures ------------------------------------------------------------------------
DGM = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
DSP = "http://schemas.microsoft.com/office/drawing/2008/diagram"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


@dataclass
class TextNodeRef:
    """A plain text value (chart series name / category label in the chart's cached data)."""
    segment_id: str
    kind: SegmentKind
    slide_no: int
    shape_key: str
    element: etree._Element


@dataclass
class ImageRef:
    slide_no: int
    blob: bytes


@dataclass
class SmartArtRef:
    """Marks a SmartArt graphic: its text is in the data part; PowerPoint displays a cached drawing part."""
    slide_no: int
    shape_key: str
    drawing: etree._Element | None


@dataclass
class Notice:
    text: str


# SmartArt parts are plain (non-XML) parts in python-pptx: parse once per presentation, write back on save.
_PARTS: dict[int, dict[str, tuple[object, etree._Element]]] = {}


def part_xml(prs: PresentationT, part: object) -> etree._Element:
    cache = _PARTS.setdefault(id(prs), {})
    key = str(part.partname)  # type: ignore[attr-defined]
    if key not in cache:
        cache[key] = (part, etree.fromstring(part.blob))  # type: ignore[attr-defined]
    return cache[key][1]


def flush_parts(prs: PresentationT) -> None:
    for part, root in _PARTS.pop(id(prs), {}).values():
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)  # type: ignore[attr-defined]


def _chart_items(shape, slide_no: int, key: str) -> Iterator[object]:
    chart = shape.chart
    frames = []
    if chart.has_title and chart.chart_title.has_text_frame:
        frames.append(("title", chart.chart_title.text_frame))
    for axis_name in ("category_axis", "value_axis"):
        try:
            axis = getattr(chart, axis_name)
        except (ValueError, KeyError):          # e.g. pie charts have no axes
            continue
        if axis.has_title and axis.axis_title.has_text_frame:
            frames.append((axis_name, axis.axis_title.text_frame))
    ck = f"{key}/chart"
    for name, tf in frames:
        for i, p in enumerate(tf.paragraphs):
            yield ParagraphRef(f"{ck}/{name}/p{i}", SegmentKind.CHART, slide_no, ck, p._p, tf, None)
    values = chart._chartSpace.xpath(".//c:strCache/c:pt/c:v")
    for i, v in enumerate(values):
        yield TextNodeRef(f"{ck}/v{i}", SegmentKind.CHART, slide_no, ck, v)
    if values:
        yield Notice(f"slide {slide_no}: chart labels translated; the chart's data sheet (Edit Data) keeps the "
                     "original labels")


def _smartart_items(shape, slide_no: int, key: str, prs: PresentationT) -> Iterator[object] | None:
    gd = shape.element.find(f".//{A}graphicData")
    if gd is None or gd.get("uri") != DGM:
        return None
    rel_ids = gd.find(f"{{{DGM}}}relIds")
    if rel_ids is None:
        return None

    def items() -> Iterator[object]:
        data = part_xml(prs, shape.part.related_part(rel_ids.get(f"{{{R_NS}}}dm")))
        drawing = None
        ext = data.find(f".//{{{DSP}}}dataModelExt")
        if ext is not None and ext.get("relId"):
            try:
                drawing = part_xml(prs, shape.part.related_part(ext.get("relId")))
            except KeyError:
                drawing = None
        sk = f"{key}/smartart"
        yield SmartArtRef(slide_no, sk, drawing)
        for i, t in enumerate(data.iter(f"{{{DGM}}}t")):
            for j, para in enumerate(t.findall(f"{A}p")):
                yield ParagraphRef(f"{sk}/n{i}/p{j}", SegmentKind.BODY, slide_no, sk, para, None, None)
    return items()


def _iter_shape(shape, slide_no: int, prefix: str, prs: PresentationT) -> Iterator[object]:
    key = f"{prefix}/{shape.shape_id}"
    st = shape.shape_type
    if st == MSO_SHAPE_TYPE.GROUP:
        for sub in shape.shapes:
            yield from _iter_shape(sub, slide_no, key, prs)
        return
    if getattr(shape, "has_table", False) and shape.has_table:
        for r, row in enumerate(shape.table.rows):
            for c, cell in enumerate(row.cells):
                if cell.is_spanned:
                    continue
                ck = f"{key}/t{r}.{c}"
                for i, p in enumerate(cell.text_frame.paragraphs):
                    yield ParagraphRef(f"{ck}/p{i}", SegmentKind.TABLE_CELL, slide_no, ck, p._p,
                                       cell.text_frame, None)
        return
    if getattr(shape, "has_chart", False) and shape.has_chart:
        yield from _chart_items(shape, slide_no, key)
        return
    if st == MSO_SHAPE_TYPE.PICTURE:
        try:
            yield ImageRef(slide_no, shape.image.blob)
        except (AttributeError, KeyError, ValueError):
            pass
        return
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        kind = SegmentKind.BODY
        if shape.is_placeholder and shape.placeholder_format.type in (
                PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            kind = SegmentKind.TITLE
        for i, p in enumerate(shape.text_frame.paragraphs):
            yield ParagraphRef(f"{key}/p{i}", kind, slide_no, key, p._p, shape.text_frame, shape)
        return
    if shape.element.tag.endswith("graphicFrame"):
        smart = _smartart_items(shape, slide_no, key, prs)
        if smart is not None:
            yield from smart
        else:
            yield f"slide {slide_no}: embedded object (not translated)"


def walk(prs: PresentationT) -> Iterator[object]:
    """Yields ParagraphRef / TextNodeRef for translatable text, ImageRef for pictures, SmartArtRef markers,
    Notice for information, and a string for content that cannot be translated."""
    for n, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            yield from _iter_shape(shape, n, f"s{n}", prs)
        if slide.has_notes_slide:
            tf = slide.notes_slide.notes_text_frame
            if tf is not None:
                for i, p in enumerate(tf.paragraphs):
                    yield ParagraphRef(f"s{n}/notes/p{i}", SegmentKind.NOTE, n, f"s{n}/notes", p._p, tf, None)

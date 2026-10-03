"""Quality control: independent checks of finished translations, file-pair checks, overflow estimation."""
from __future__ import annotations

from pathlib import Path

import pymupdf
import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from sbt import languages, quality
from sbt.app import checks
from sbt.domain.models import DocumentModel, Segment, SegmentKind
from sbt.quality import compare
from sbt.renderers.pptx_renderer import PptxRenderer
from sbt.renderers.text_fit import _inherited_sz, estimate_height
from sbt.terminology.glossary import Glossary, Term
from sbt.terminology.term_sheet import DocTerm, TermSheet
from sbt.ui import pages

G = Glossary([Term("Building Management System", "ビル管理システム", "en", "ja"),
              Term("BACnet", "BACnet", "en", "ja", do_not_translate=True)])


def _doc(*pairs: tuple[str, str | None]) -> DocumentModel:
    segs = [Segment(f"s{i + 1}/1/p0", SegmentKind.BODY, i + 1, f"s{i + 1}/1", src, translation=tgt)
            for i, (src, tgt) in enumerate(pairs)]
    return DocumentModel("x.pptx", "pptx", len(segs), segs)


def _codes(report: quality.QcReport, sid: str | None = None) -> set[str]:
    return {f.code for f in report.findings if sid is None or sid in f.segments}


def test_clean_translation_has_no_findings() -> None:
    doc = _doc(("The Building Management System monitors the chillers.", "ビル管理システムは冷凍機を監視します。"),
               ("BACnet gateway at 192.168.1.100", "BACnetゲートウェイ（192.168.1.100）"),
               ("Supply air temperature: 24 °C", "給気温度：24 °C"))
    assert quality.check(doc, "en", "ja", G).findings == []


def test_missing_and_partial_translation() -> None:
    doc = _doc(("The controller reports alarms.", "The controller reports alarms."),
               ("Edge devices send telemetry to the cloud.", "エッジ devices はテレメトリを cloud に送信します。"),
               ("The gateway publishes data every minute.", "ゲートウェイは毎分データを"),
               ("Occupancy sensors and CO2 sensors", None))
    rep = quality.check(doc, "en", "ja", G)
    assert "not_translated" in _codes(rep, "s1/1/p0") and "not_translated" in _codes(rep, "s4/1/p0")
    partial = next(f for f in rep.findings if f.code == "partly_translated")
    assert partial.detail["left"] == ["devices", "cloud"]
    assert "possible_omission" in _codes(rep, "s3/1/p0")          # ends mid-sentence


def test_dangling_word_means_cut_off() -> None:
    rep = quality.check(_doc(("在室センサーとCO2センサー", "Occupancy sensor and")), "ja", "en", Glossary([]))
    assert "possible_omission" in _codes(rep)


def test_identifiers_numbers_and_units() -> None:
    doc = _doc(("Gateway 192.168.1.100, firmware v3.2.1", "ゲートウェイ 192.168.1.10、ファームウェア v3.2.1"),
               ("Cooling setpoint raised from 24 °C to 26 °C", "冷房設定値を24 °Fから26 °Cに引き上げ"),
               ("Peak demand is 1,250 kW", "ピーク需要は1,520 kW"),
               ("See https://bms.example.com", "https://bms.example.com と https://evil.example.org を参照"))
    rep = quality.check(doc, "en", "ja", Glossary([]))
    changed = next(f for f in rep.findings if f.code == "token_changed")
    assert changed.detail == {"from": "192.168.1.100", "to": "192.168.1.10"}
    assert "number_changed" in _codes(rep, "s2/1/p0")              # one of two °C became °F
    assert "number_changed" in _codes(rep, "s3/1/p0")
    assert "token_added" in _codes(rep, "s4/1/p0")


def test_units_written_out_are_not_flagged() -> None:
    doc = _doc(("Energy consumption: 1,250 MWh/year", "能源消耗：每年1,250兆瓦时"))
    assert quality.check(doc, "en", "zh", Glossary([])).findings == []


def test_terminology_glossary_variants_and_repeated_text() -> None:
    doc = _doc(("Building Management System overview", "建物管理システムの概要"),
               ("The chilled water plant runs at night.", "冷水プラントは夜間に運転します。"),
               ("Chilled water plant efficiency", "冷水プラントの効率"),
               ("The chilled water plant has three chillers.", "冷水設備には冷凍機が3台あります。"),
               ("Alarm list", "アラーム一覧"), ("Alarm list", "警報リスト"))
    sheet = TermSheet("en", "ja", [DocTerm("chilled water plant", {"s2/1/p0", "s3/1/p0", "s4/1/p0"}, [],
                                           ["冷水プラント", "冷水設備"], "冷水プラント")])
    rep = quality.check(doc, "en", "ja", G, sheet=sheet)
    glossary = next(f for f in rep.findings if f.code == "glossary_term")
    assert glossary.segments == ["s1/1/p0"] and glossary.detail["approved"] == "ビル管理システム"
    variants = next(f for f in rep.findings if f.code == "term_variants")
    assert variants.segments == ["s4/1/p0"] and variants.severity.value == "warning"
    assert "same_source_differs" in _codes(rep, "s6/1/p0")
    assert all(f.describe() for f in rep.findings)
    assert "冷水" not in str(rep.public())                          # report.json part has no document text


def _deck(path: Path, texts: list[str]) -> Path:
    prs = Presentation()
    for text in texts:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(0.6))
        box.text_frame.text = text
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    prs.save(str(path))
    return path


def test_written_file_is_read_back(tmp_path: Path) -> None:
    deck = _deck(tmp_path / "a.pptx", ["Chiller plant status", "Gateway configuration"])
    from sbt.parsers.pptx_parser import PptxParser
    doc = PptxParser(ocr_images=False).parse(deck)
    doc.segments[0].translation, doc.segments[1].translation = "冷水プラントの状態", "ゲートウェイの設定"
    out = tmp_path / "a_JA.pptx"
    written = PptxParser(ocr_images=False).parse(deck)            # the writer "forgets" the second paragraph
    written.segments[0].translation = doc.segments[0].translation
    PptxRenderer("en", "ja").render(written, out)
    rep = quality.check(doc, "en", "ja", Glossary([]), output=out, original=deck)
    assert _codes(rep, doc.segments[1].id) == {"not_written"} and not _codes(rep, doc.segments[0].id)


def test_pdf_paragraph_must_be_written_where_it_was(tmp_path: Path) -> None:
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 100), "Protocol")
    page.insert_text((72, 300), "プロトコル", fontname="japan")       # the same word elsewhere on the page
    out = tmp_path / "out.pdf"
    pdf.save(out)
    seg = Segment("p1/r0", SegmentKind.BODY, 1, "p1/r0", "Protocol", translation="プロトコル",
                  geometry={"page": 0, "rect": (72, 88, 130, 104)})
    rep = quality.check(DocumentModel("in.pdf", "pdf", 1, [seg]), "en", "ja", Glossary([]), output=out)
    assert {"not_written", "original_left"} <= _codes(rep)


def test_subtitle_uses_body_size_and_spacing_counts() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])               # title slide: title + subtitle
    title, subtitle = slide.shapes.title, slide.placeholders[1]
    assert _inherited_sz(subtitle, 0) < _inherited_sz(title, 0)      # "SUBTITLE" is not a title
    body = prs.slides.add_slide(prs.slide_layouts[1]).placeholders[1]
    body.text_frame.text = "First point"
    one = estimate_height(body, languages.get("en"))[0]
    body.text_frame.add_paragraph().text = "Second point"
    two = estimate_height(body, languages.get("en"))[0]
    assert two > 2 * one                                             # space before the second paragraph


def test_check_a_translation_made_elsewhere(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("SBT_DATA_DIR", str(tmp_path / "data"))
    original = _deck(tmp_path / "spec.pptx", ["BACnet gateway at 192.168.1.100", "Building Management System"])
    translated = _deck(tmp_path / "spec_agency.pptx", ["BACnetゲートウェイ（192.168.1.10）", "建物管理システム"])
    doc = compare.load_pair(original, translated, "en")
    assert [s.translation for s in doc.segments] == ["BACnetゲートウェイ（192.168.1.10）", "建物管理システム"]
    from sbt.settings import Settings
    job = checks.run_check(checks.CheckSpec(original, translated, "en", "ja", use_model=False), Settings())
    assert {"token_changed", "glossary_term"} <= _codes(job.quality)
    assert job.report["matched"] == 2 and not job.term_note          # model not requested: no note
    csv_text = job.write_csv(tmp_path / "checks.csv").read_text(encoding="utf-8-sig")
    assert "192.168.1.10" in csv_text
    with pytest.raises(Exception, match="different number"):
        compare.load_pair(original, _deck(tmp_path / "short.pptx", ["x"]), "en")


def test_pdf_pages_side_by_side(tmp_path: Path) -> None:
    pdf = pymupdf.open()
    pdf.new_page().insert_text((72, 100), "Chiller")
    pdf.save(tmp_path / "a.pdf")
    pdf.save(tmp_path / "b.pdf")
    seg = Segment("p1/r0", SegmentKind.BODY, 1, "p1/r0", "Chiller", translation="冷凍機",
                  geometry={"page": 0, "rect": (72, 88, 120, 104)})
    doc = DocumentModel(str(tmp_path / "a.pdf"), "pdf", 1, [seg])
    report = quality.check(doc, "en", "ja", Glossary([]), output=tmp_path / "b.pdf")
    view = pages.page_view(doc, report, tmp_path / "a.pdf", tmp_path / "b.pdf", 1)
    assert view["available"] and str(view["original"]).startswith("data:image/png;base64,")
    assert view["boxes"] and view["boxes"][0]["severity"] == "error"   # not written: outlined

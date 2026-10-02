"""Desktop UI back end: the API the window calls, background jobs, review edits, settings screen."""
from __future__ import annotations

import time
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from sbt.app.jobs import JobCancelled, JobSpec, TranslationJob
from sbt.settings import Settings
from sbt.terminology.glossary import Glossary
from sbt.ui import review, settings_store
from sbt.ui.api import Api
from sbt.ui.runner import JobRunner, UiReporter
from tests.unit.test_phase2 import FakeEngine

TABLE = {"Chiller plant": "冷水プラント", "status": "状態", "Open the": "", "BACnet": "BACnet",
         "gateway": "ゲートウェイを開く", " ": ""}


def _deck(path: Path) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tf = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2)).text_frame
    tf.text = "Chiller plant status"
    p = tf.add_paragraph()
    for text, bold in (("Open the ", False), ("BACnet", True), (" gateway", False)):
        run = p.add_run()
        run.text, run.font.bold = text, bold
    prs.save(str(path))
    return path


def _fake_run(spec: JobSpec, settings: Settings, use_memory: bool = True,
              reporter: UiReporter | None = None) -> TranslationJob:
    job = TranslationJob(spec, Glossary([]), None, reporter)
    if reporter:
        reporter.stage("load")
    job.translate(FakeEngine(TABLE))
    job.finish()
    return job


def _wait(runner: JobRunner) -> None:
    assert runner._thread is not None
    runner._thread.join(timeout=60)
    assert not runner.busy


@pytest.fixture()
def data(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    monkeypatch.setenv("SBT_DATA_DIR", str(tmp_path / "data"))
    return tmp_path


def test_tags_are_shown_as_numbers_and_converted_back() -> None:
    assert review.to_display("<g1>BACnet</g1> gateway") == "[1]BACnet[/1] gateway"
    assert review.from_display("[1]BACnet[/1] and [3]x[/3]", {1}) == "<g1>BACnet</g1> and [3]x[/3]"


def test_translate_review_edit_and_rebuild(data: Path) -> None:
    deck = _deck(data / "plant.pptx")
    original = deck.read_bytes()
    api = Api(JobRunner(run=_fake_run))
    info = api.inspect_file(str(deck))
    assert info["type"] == "PPTX" and info["detected"] == "en" and info["target"] == "ja"
    assert info["paragraphs"] == 2

    with pytest.raises(ValueError):
        api.start_translation({"path": str(deck), "source": "en", "target": "en"})
    out = api.start_translation({"path": str(deck), "source": "en", "target": "ja", "memory": False})
    assert out["output"].endswith("plant_JA.pptx")
    _wait(api._runner)
    status = api.job_status()
    assert status["status"] == "done", status["error"]
    assert status["stages_done"][:2] == ["read", "load"] and "write" in status["stages_done"]
    result = status["result"]
    assert result["segments"] == 2 and result["translated_pct"] == 100.0
    assert Path(result["output"]).exists()

    items = api.review_items()
    tagged = next(i for i in items if i["has_tags"])
    assert "[1]" in tagged["source"] or "[2]" in tagged["source"]

    # Formatting markers that do not match the original: saved anyway, without markers, with a warning.
    warnings = api.check_edit(tagged["id"], "[9]BACnet ゲートウェイを開く")
    assert any("Formatting markers" in w for w in warnings)
    plain = next(i for i in items if not i["has_tags"])
    saved = api.save_edit(plain["id"], "冷水プラントの運転状態")
    assert saved["translation"] == "冷水プラントの運転状態"
    assert next(i for i in api.review_items() if i["id"] == plain["id"])["status"] == "edited"

    rebuilt = api.rebuild()
    assert rebuilt["edited"] == 1
    texts = [sh.text_frame.text for sh in Presentation(result["output"]).slides[0].shapes if sh.has_text_frame]
    assert "冷水プラントの運転状態" in texts[0]
    assert deck.read_bytes() == original


def test_cancel_stops_the_job(data: Path) -> None:
    def slow(spec: JobSpec, settings: Settings, use_memory: bool = True,
             reporter: UiReporter | None = None) -> TranslationJob:
        assert reporter is not None
        reporter.stage("translate")
        for i in range(1000):
            reporter.progress(i, 1000)
            time.sleep(0.01)
        raise AssertionError("not cancelled")

    runner = JobRunner(run=slow)
    runner.start(JobSpec(Path("a.pptx"), Path("b.pptx"), "en", "ja"), Settings(), use_memory=False)
    with pytest.raises(RuntimeError):
        runner.start(JobSpec(Path("a.pptx"), Path("b.pptx"), "en", "ja"), Settings(), use_memory=False)
    time.sleep(0.05)
    runner.cancel()
    _wait(runner)
    assert runner.state.status == "cancelled"


@pytest.mark.parametrize(("error", "expected"), [
    (PermissionError("x"), "Close it and try again"),
    (FileNotFoundError("x"), "Open Models"),
    (RuntimeError("Confidential slide text"), "RuntimeError"),
])
def test_errors_become_plain_messages_without_document_text(error: Exception, expected: str) -> None:
    def failing(*args: object, **kwargs: object) -> TranslationJob:
        raise error

    runner = JobRunner(run=failing)
    runner.start(JobSpec(Path("a.pptx"), Path("b.pptx"), "en", "ja"), Settings(), use_memory=False)
    _wait(runner)
    assert runner.state.status == "failed"
    assert expected in runner.state.error and "Confidential" not in runner.state.error


def test_reporter_raises_when_cancelled() -> None:
    from threading import Event

    from sbt.ui.runner import JobState
    state, cancel = JobState(), Event()
    rep = UiReporter(state, cancel)
    rep.stage("read")
    rep.stage("translate")
    assert state.stages_done == ["read"] and state.stage == "translate"
    cancel.set()
    with pytest.raises(JobCancelled):
        rep.progress(1, 2)


def test_settings_screen_saves_and_validates(data: Path) -> None:
    saved = settings_store.save({"theme": "dark", "gpu": "cpu", "use_memory": False})
    assert saved["theme"] == "dark"
    now = settings_store.current()
    assert now["theme"] == "dark" and now["gpu"] == "cpu" and now["use_memory"] is False
    with pytest.raises(ValueError):
        settings_store.save({"gpu": "quantum"})
    with pytest.raises(ValueError):
        settings_store.save({"server_url": "https://example.com"})
    assert settings_store.current()["gpu"] == "cpu"


def test_api_exposes_no_public_attributes() -> None:
    """pywebview exposes public attributes to JavaScript; only methods may be public."""
    api = Api(JobRunner(run=_fake_run))
    assert [k for k in vars(api) if not k.startswith("_")] == []

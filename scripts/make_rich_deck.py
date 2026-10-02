"""Public test decks with a chart and pictures (rich_en.pptx, rich_ja.pptx).

SmartArt cannot be created with python-pptx (and PowerPoint on the dev PC is unlicensed, so it cannot create it
either); SmartArt is tested with Apache POI's public samples in eval/testset/external/.
"""
from __future__ import annotations

import io
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent.parent / "eval" / "testset"
CONTENT = {
    "en": {"sa_title": "Data Flow", "nodes": ["Field devices", "Controllers", "Gateway", "Cloud platform"],
           "chart_title": "Energy Consumption by System", "cats": ["Chiller", "Air handling units", "Lighting",
                                                                   "Other loads"],
           "series": ["Energy consumption (MWh)", "Target (MWh)"], "x": "System", "y": "MWh per year",
           "pic_title": "Commissioning Photos", "pic_text": ["Supply air temperature", "Setpoint 18 °C"],
           "font": "arial.ttf"},
    "ja": {"sa_title": "データの流れ", "nodes": ["フィールド機器", "コントローラー", "ゲートウェイ", "クラウドプラットフォーム"],
           "chart_title": "システム別エネルギー消費量", "cats": ["冷凍機", "空調機", "照明", "その他の負荷"],
           "series": ["エネルギー消費量（MWh）", "目標値（MWh）"], "x": "システム", "y": "年間MWh",
           "pic_title": "試運転写真", "pic_text": ["給気温度", "設定値 18℃"], "font": "meiryo.ttc"},
}


def _picture(lines: list[str], font: str, with_text: bool) -> bytes:
    img = Image.new("RGB", (900, 420), (232, 238, 245))
    d = ImageDraw.Draw(img)
    for i in range(0, 900, 30):                                   # some non-text structure
        d.line([(i, 0), (i + 200, 420)], fill=(210, 220, 232), width=8)
    if with_text:
        d.rectangle([150, 120, 750, 300], fill=(30, 30, 30))
        f = ImageFont.truetype(rf"C:\Windows\Fonts\{font}", 56)
        for k, line in enumerate(lines):
            d.text((190, 140 + k * 76), line, fill=(120, 255, 120), font=f)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def build(lang: str) -> None:
    c = CONTENT[lang]
    path = OUT / f"rich_{lang}.pptx"
    prs = Presentation()
    layout = prs.slide_layouts[5]                                # title only
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = c["chart_title"]
    data = CategoryChartData()
    data.categories = c["cats"]
    data.add_series(c["series"][0], (420, 260, 140, 90))
    data.add_series(c["series"][1], (380, 240, 120, 80))
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.7), Inches(1.6), Inches(8.6),
                                   Inches(4.6), data).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = c["chart_title"]
    for axis, title in ((chart.category_axis, c["x"]), (chart.value_axis, c["y"])):
        axis.has_title = True
        axis.axis_title.text_frame.text = title
        axis.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    chart.has_legend = True
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = c["pic_title"]
    for k, with_text in enumerate((True, False)):
        slide.shapes.add_picture(io.BytesIO(_picture(c["pic_text"], c["font"], with_text)),
                                 Inches(0.4 + k * 4.7), Inches(2.0), Inches(4.4))
    prs.save(str(path))
    print("wrote", path.name)


if __name__ == "__main__":
    for lang in CONTENT:
        build(lang)

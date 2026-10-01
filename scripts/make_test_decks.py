"""Generate public, non-confidential test decks (EN and JA) for the Phase 1 evaluation.

Covers: titles, multi-level bullets, mixed bold/italic/colour runs, a table, grouped shapes, speaker notes,
numbers/units/percentages/dates, URLs, e-mails, IP addresses, part numbers, acronyms and long sentences.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent.parent / "eval" / "testset"
BLUE = RGBColor(0x1F, 0x4E, 0x9A)
RED = RGBColor(0xC0, 0x1F, 0x1F)

# Rich text: list of (text, style) where style is a subset of "b", "i", "blue", "red".
Rich = list[tuple[str, str]]


def add_rich(paragraph, parts: Rich) -> None:
    for text, style in parts:
        run = paragraph.add_run()
        run.text = text
        run.font.bold = "b" in style or None
        run.font.italic = "i" in style or None
        if "blue" in style:
            run.font.color.rgb = BLUE
        if "red" in style:
            run.font.color.rgb = RED


def bullets_slide(prs, title: str, items: list[tuple[int, Rich | str]], notes: str = "") -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    tf = slide.placeholders[1].text_frame
    for i, (level, content) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = level
        add_rich(p, [(content, "")] if isinstance(content, str) else content)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def table_slide(prs, title: str, rows: list[list[str]]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(0.5), Inches(1.6), Inches(9), Inches(4))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            shape.table.cell(r, c).text = value


def diagram_slide(prs, title: str, boxes: list[str], caption: str) -> None:
    """Grouped shapes with text: architecture-diagram style."""
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    group = slide.shapes.add_group_shape()
    for i, label in enumerate(boxes):
        box = group.shapes.add_shape(1, Inches(0.4 + i * 2.35), Inches(2.2), Inches(2.1), Inches(1.2))
        box.text_frame.text = label
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(9), Inches(1.0))
    tb.text_frame.word_wrap = True
    tb.text_frame.text = caption
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(18)


def build_en() -> Presentation:
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Smart Building Platform Overview"
    s.placeholders[1].text = "Building Management System (BMS) Integration Proposal — 15 March 2026"

    bullets_slide(prs, "Building Management System", [
        (0, [("The ", ""), ("Building Management System", "b"), (" monitors and controls HVAC equipment.", "")]),
        (1, "Air Handling Unit (AHU) and Fan Coil Unit (FCU) control"),
        (1, "Variable Air Volume (VAV) terminal control"),
        (0, [("Alarm", "b red"), (" and ", ""), ("equipment status", "i"), (" monitoring 24/7", "")]),
        (0, "Energy Management System (EMS) for energy consumption analysis"),
    ], notes="The building management system collects data from more than 2,000 points and sends "
             "alarms to the operator within 5 seconds.")

    bullets_slide(prs, "Sensors and Field Devices", [
        (0, "Occupancy sensor"),
        (0, "CO2 sensor (0–2,000 ppm)"),
        (0, "Temperature sensor: ±0.2 °C accuracy"),
        (0, "Humidity: 30–70 %RH"),
        (0, [("Model ", ""), ("FX-PCG2611-0", "b"), (" controller, firmware v3.2.1", "")]),
    ])

    diagram_slide(prs, "System Architecture",
                  ["Edge device", "Gateway", "Cloud platform", "Dashboard"],
                  "Field controllers communicate over BACnet/IP and Modbus TCP; the gateway publishes "
                  "telemetry to the cloud platform via MQTT over TLS (port 8883).")

    table_slide(prs, "Protocol Comparison", [
        ["Protocol", "Transport", "Typical use", "Default port"],
        ["BACnet/IP", "UDP", "HVAC controllers", "47808"],
        ["Modbus TCP", "TCP", "Power meters, chillers", "502"],
        ["MQTT", "TCP", "IoT telemetry to cloud", "1883 / 8883"],
        ["KNX", "Twisted pair / IP", "Lighting and blinds", "3671"],
    ])

    bullets_slide(prs, "Energy Efficiency and Demand Response", [
        (0, "Peak demand reduced by 18% in summer 2025"),
        (0, "Energy consumption: 1,250 MWh/year → 1,020 MWh/year"),
        (0, [("Demand response", "b"), (" signals are received from the utility via ", ""), ("OpenADR 2.0b", "blue")]),
        (0, "Setpoint adjustment: cooling setpoint raised from 24 °C to 26 °C during peak hours"),
        (0, "Power consumption is measured every 15 minutes"),
    ], notes="Energy saving control is enabled only when the occupancy rate is below 30 percent "
             "and the outdoor air temperature is above 32 °C.")

    bullets_slide(prs, "Network and Access", [
        (0, "BMS server: 192.168.1.100 (VLAN 20)"),
        (0, "Web dashboard: https://bms.example.com/dashboard"),
        (0, "Support: support@example.com"),
        (0, r"Configuration file: C:\Program Files\BMS\config\points.json"),
        (0, "REST API over HTTPS; payloads in JSON, legacy exports in XML and SQL"),
    ])

    bullets_slide(prs, "Operation and Maintenance", [
        (0, "When a chiller alarm occurs, the system automatically notifies the maintenance team by e-mail "
            "and creates a work order, and if the alarm is not acknowledged within 30 minutes, it is "
            "escalated to the facility manager."),
        (0, "Scheduled maintenance of the AHU filters is performed every 3 months."),
        (0, "The PLC and SCADA system remain in service during the migration period."),
    ])

    bullets_slide(prs, "Building Automation Roadmap", [
        (0, "Phase 1 (Q2 2026): Building automation system upgrade"),
        (0, "Phase 2 (Q4 2026): IoT sensors and edge devices"),
        (0, "Phase 3 (2027): AI-based fault detection and diagnostics (FDD)"),
    ])
    return prs


def build_ja() -> Presentation:
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "スマートビルディングプラットフォーム概要"
    s.placeholders[1].text = "ビル管理システム（BMS）統合提案 — 2026年3月15日"

    bullets_slide(prs, "ビル管理システム", [
        (0, [("ビル管理システム", "b"), ("は空調設備を監視・制御します。", "")]),
        (1, "空調機（AHU）およびファンコイルユニット（FCU）の制御"),
        (1, "可変風量（VAV）ターミナル制御"),
        (0, [("警報", "b red"), ("および", ""), ("機器状態", "i"), ("の24時間監視", "")]),
        (0, "エネルギー消費量分析のためのエネルギー管理システム（EMS）"),
    ], notes="ビル管理システムは2,000点以上のポイントからデータを収集し、5秒以内にオペレーターへ警報を送信します。")

    bullets_slide(prs, "センサーおよびフィールド機器", [
        (0, "在室センサー"),
        (0, "CO2センサー（0～2,000 ppm）"),
        (0, "温度センサー：精度±0.2 °C"),
        (0, "湿度：30～70 %RH"),
        (0, [("型式 ", ""), ("FX-PCG2611-0", "b"), (" コントローラー、ファームウェア v3.2.1", "")]),
    ])

    diagram_slide(prs, "システム構成",
                  ["エッジデバイス", "ゲートウェイ", "クラウドプラットフォーム", "ダッシュボード"],
                  "フィールドコントローラーはBACnet/IPおよびModbus TCPで通信し、ゲートウェイはMQTT over TLS"
                  "（ポート8883）でクラウドプラットフォームへテレメトリを送信します。")

    table_slide(prs, "プロトコル比較", [
        ["プロトコル", "トランスポート", "主な用途", "デフォルトポート"],
        ["BACnet/IP", "UDP", "空調コントローラー", "47808"],
        ["Modbus TCP", "TCP", "電力量計、冷凍機", "502"],
        ["MQTT", "TCP", "クラウドへのIoTテレメトリ", "1883 / 8883"],
        ["KNX", "ツイストペア / IP", "照明およびブラインド", "3671"],
    ])

    bullets_slide(prs, "省エネルギーとデマンドレスポンス", [
        (0, "2025年夏にピーク需要を18%削減"),
        (0, "エネルギー消費量：1,250 MWh/年 → 1,020 MWh/年"),
        (0, [("デマンドレスポンス", "b"), ("信号は", ""), ("OpenADR 2.0b", "blue"), ("経由で電力会社から受信", "")]),
        (0, "設定値調整：ピーク時間帯は冷房設定値を24 °Cから26 °Cに変更"),
        (0, "消費電力は15分ごとに計測"),
    ], notes="省エネルギー制御は、在室率が30%未満かつ外気温度が32 °Cを超える場合にのみ有効になります。")

    bullets_slide(prs, "ネットワークとアクセス", [
        (0, "BMSサーバー：192.168.1.100（VLAN 20）"),
        (0, "Webダッシュボード：https://bms.example.com/dashboard"),
        (0, "サポート：support@example.com"),
        (0, r"設定ファイル：C:\Program Files\BMS\config\points.json"),
        (0, "HTTPS経由のREST API。ペイロードはJSON、旧形式のエクスポートはXMLおよびSQL"),
    ])

    bullets_slide(prs, "運用・保守", [
        (0, "冷凍機の警報が発生すると、システムは保守チームに自動的にメールで通知して作業指示を作成し、"
            "30分以内に警報が確認されない場合は施設管理者にエスカレーションします。"),
        (0, "空調機フィルターの定期保守は3か月ごとに実施します。"),
        (0, "移行期間中もPLCおよびSCADAシステムは稼働を継続します。"),
    ])
    return prs


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    build_en().save(OUT / "sample_en.pptx")
    build_ja().save(OUT / "sample_ja.pptx")
    print(f"Wrote {OUT / 'sample_en.pptx'} and {OUT / 'sample_ja.pptx'}")


if __name__ == "__main__":
    main()
